from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "StayHub_Alcance_Funcional_y_Correcciones_Manual.docx"
NAVY = "25344F"
PALE = "F2F5FA"
GRID = "D9D9D9"


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.find(qn("w:tcBorders"))
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), GRID)


def cell_margin(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text, *, bold=False, white=False, center=False, size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
    if white:
        run.font.color.rgb = RGBColor(255, 255, 255)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    cell_margin(cell)
    borders(cell)


def add_table(doc, headers, rows, widths=None, centered_columns=()):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.rows[0]._tr.get_or_add_trPr()
    repeat_header(table.rows[0])
    for index, header in enumerate(headers):
        shade(table.rows[0].cells[index], NAVY)
        set_cell_text(table.rows[0].cells[index], header, bold=True, white=True, center=True, size=9)
        if widths:
            table.rows[0].cells[index].width = Inches(widths[index])
    for row_index, row_data in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(row_data):
            if row_index % 2:
                shade(cells[col_index], PALE)
            set_cell_text(cells[col_index], value, center=col_index in centered_columns)
            if widths:
                cells[col_index].width = Inches(widths[col_index])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    run.font.size = Pt(9)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.12

for name, size, before, after in (("Title", 26, 0, 16), ("Heading 1", 17, 14, 7), ("Heading 2", 13, 11, 5)):
    style = styles[name]
    style.font.name = "Aptos Display"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = name != "Title"
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

footer = section.footer.paragraphs[0]
add_page_number(footer)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Alcance funcional vigente de StayHub")
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(22)
run = subtitle.add_run("Correcciones para el manual y la documentación técnica")
run.font.size = Pt(15)
run.font.color.rgb = RGBColor(68, 71, 77)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.paragraph_format.space_after = Pt(22)
meta.add_run("Desarrollo de Aplicaciones II\nGrupo 2\n30 de octubre de 2026\nKAN-102 y KAN-103")

intro = doc.add_paragraph()
intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
intro.add_run(
    "Este documento establece el comportamiento funcional vigente de la rama development y "
    "corrige diferencias detectadas entre el README, el manual de usuario y la documentación de "
    "componentes. La versión actual incluye los flujos principales de reserva y administración, "
    "pero conserva simuladores para pagos, OTAs, PMS, SMS y PUSH. Favoritos, vencimiento automático "
    "de holds y la integración completa de Overbooking no están implementados."
)

doc.add_page_break()

doc.add_heading("Conclusión de la revisión", level=1)
doc.add_paragraph(
    "La documentación ya puede describir un único alcance técnico, siempre que diferencie los dos "
    "modos de ejecución y no presente los simuladores como proveedores productivos. Un pago aprobado "
    "confirma las reservas antes de responder. La resolución de Overbooking también ocurre dentro de "
    "la llamada HTTP; sólo la notificación posterior se procesa mediante JMS."
)

doc.add_heading("Direcciones por entorno", level=1)
add_table(
    doc,
    ["Entorno", "Aplicación", "API", "WSDL SOAP"],
    [
        ["Docker Compose", "https://localhost:8443/StayHub/index.html", "https://localhost:8443/StayHub/api", "https://localhost:8443/StayHub/soap/integracion-pms?wsdl"],
        ["WildFly manual endurecido", "https://localhost:8080/StayHub/index.html", "https://localhost:8080/StayHub/api", "https://localhost:8080/StayHub/soap/integracion-pms?wsdl"],
    ],
    widths=[1.25, 1.9, 1.55, 2.2],
)
doc.add_paragraph(
    "El puerto 8080 funciona como HTTPS únicamente después de aplicar la configuración manual. Una "
    "instalación de WildFly existente puede conservar otro puerto; debe verificarse antes de la demo."
)

doc.add_heading("Alcance implementado", level=1)
implemented = [
    ["Catálogo", "Implementado", "Hoteles, tipos y habitaciones con bajas lógicas."],
    ["Inventario y tarifas", "Implementado", "Cupo y precio por noche y holds persistidos."],
    ["Reserva y carrito", "Implementado", "El carrito conserva la selección por sesión y crea reservas pendientes."],
    ["Reservas OTA", "Con simuladores", "REST delega en Reservas e identifica cada operación por canal y referencia."],
    ["Pago", "Con simulador", "El pago aprobado confirma las reservas y luego publica eventos JMS."],
    ["Sincronización OTA y PMS", "Con simuladores", "La API encola; un MDB publica por REST o SOAP."],
    ["Overbooking", "Limitado", "Resuelve y persiste en forma síncrona; la notificación se encola."],
    ["Notificaciones", "Mixto", "EMAIL puede ser real; SMS y PUSH permanecen simulados."],
    ["Usuarios y sesiones", "Implementado", "Registro, login, perfil, recuperación, roles y propiedad."],
]
add_table(doc, ["Función", "Estado", "Comportamiento"], implemented, widths=[1.55, 1.15, 4.2], centered_columns=(1,))

doc.add_heading("Funciones simuladas", level=1)
add_bullets(doc, [
    "La pasarela genera una referencia MOCK y no realiza cobros bancarios.",
    "Las OTAs se ejercitan contra endpoints controlados; no hay cuentas productivas conectadas.",
    "El PMS legado se demuestra con pms-legado-mock y un contrato SOAP.",
    "SMS y PUSH registran el procesamiento, pero no entregan mensajes externos.",
    "EMAIL realiza una entrega real sólo cuando SMTP está configurado fuera del repositorio.",
])

doc.add_heading("Funciones pendientes y fuera de alcance", level=1)
pending = [
    ["Favoritos", "Fuera de alcance", "El corazón es decorativo; no hay endpoint ni persistencia."],
    ["Vencimiento automático de holds", "Pendiente", "El cupo permanece retenido hasta confirmar o liberar el hold."],
    ["Detección automática de overbooking", "Pendiente", "El administrador debe informar el conflicto por REST."],
    ["Actualizar la reserva reubicada", "Pendiente", "No modifica la reserva original ni libera su hold anterior."],
    ["Cancelación anticipada", "Pendiente", "La estrategia existe en el enum, pero no se ejecuta."],
    ["SMS y PUSH reales", "Fuera de alcance", "No existe integración con proveedores externos."],
    ["Pasarela bancaria", "Fuera de alcance", "No existe captura, devolución ni conciliación real."],
    ["OTAs productivas", "Fuera de alcance", "La demostración depende de mocks o endpoints controlados."],
    ["Habitación física", "Fuera de alcance", "La reserva se realiza por tipo de habitación."],
]
add_table(doc, ["Elemento", "Clasificación", "Consecuencia"], pending, widths=[1.75, 1.35, 3.8], centered_columns=(1,))
doc.add_paragraph(
    "Toda reducción posterior al cierre de diseño necesita aprobación escrita de la cátedra. Esta "
    "clasificación técnica no reemplaza esa aprobación; hasta obtenerla debe figurar como pendiente "
    "de validación de alcance en Jira y en el informe final."
)

doc.add_heading("Correcciones para el manual de usuario", level=1)
corrections = [
    ["Acceso", "Separar Docker 8443 de WildFly manual 8080 y usar siempre HTTPS."],
    ["Pago", "Indicar que el pago aprobado confirma la reserva antes de responder; JMS envía correos después."],
    ["Overbooking", "Describir la resolución como síncrona y la notificación como asíncrona."],
    ["Reubicación", "Aclarar que se confirma un hold alternativo sin modificar la reserva original ni liberar el anterior."],
    ["Favoritos", "Indicar que no forma parte de la versión actual aunque se muestre un corazón decorativo."],
    ["Holds", "Mantener la advertencia de que no vencen automáticamente."],
    ["Integraciones", "Identificar como simuladas la pasarela, las OTAs, el PMS, SMS y PUSH; EMAIL puede ser real."],
]
add_table(doc, ["Sección", "Texto funcional que debe conservarse"], corrections, widths=[1.45, 5.45])

doc.add_heading("Configuración externa necesaria", level=1)
add_bullets(doc, [
    "JPA utiliza StayHubPU y el datasource java:/PostgresDS.",
    "WildFly debe iniciarse con standalone-full.xml para disponer de JMS.",
    "Cada OTA requiere stayhub.ota.<canal>.url, token opcional y hosts permitidos.",
    "El PMS saliente requiere stayhub.pms.wsdl y un host permitido.",
    "EMAIL requiere propiedades SMTP privadas; las credenciales no se versionan.",
    "Cada instalación manual necesita su propio certificado, datasource y configuración de WildFly.",
])

doc.add_heading("Evidencia comprobada en el código", level=1)
evidence = [
    ["Pago y reserva", "ServicioDePagosImpl.cobrarYConfirmar confirma cada reserva antes de publicar EventoPagoAprobado."],
    ["Overbooking", "ServicioDeOverbookingImpl resuelve y persiste antes de publicar la solicitud JMS."],
    ["SMS y PUSH", "SmsNotificador y PushNotificador registran una operación simulada."],
    ["Favoritos", "No existe recurso, entidad ni repositorio de Favoritos."],
    ["Holds", "No existe temporizador EJB ni tarea programada que los expire."],
    ["Puertos", "compose.yaml publica 8443 y config/https-8080.cli configura el modo manual en 8080."],
]
add_table(doc, ["Tema", "Evidencia"], evidence, widths=[1.4, 5.5])

doc.add_heading("Decisiones pendientes de validación externa", level=1)
doc.add_paragraph(
    "El equipo debe presentar a la cátedra la exclusión de Favoritos, el vencimiento automático de "
    "holds, los proveedores reales y la integración completa de Overbooking. La aprobación debe "
    "quedar enlazada en KAN-103. Sin esa evidencia, el trabajo técnico y documental queda completo, "
    "pero el criterio de acuerdo escrito permanece abierto."
)

doc.core_properties.title = "Alcance funcional vigente de StayHub"
doc.core_properties.subject = "Correcciones para el manual y la documentación técnica"
doc.core_properties.author = "Grupo 2 StayHub"
doc.core_properties.keywords = "StayHub, alcance funcional, KAN-102, KAN-103"
doc.save(OUTPUT)
print(OUTPUT)
